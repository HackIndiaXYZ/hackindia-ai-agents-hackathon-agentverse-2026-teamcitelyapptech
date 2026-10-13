"""Generate fictional originals and matching citation fixtures. No external data/services."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
root=Path(__file__).resolve().parents[1];out=root/'public/demo-files';out.mkdir(exist_ok=True)
data=[]
def record(id,name,pages,answers,initial=True):
 data.append(dict(id=id,filename=name,pages=[dict(page=i+1,text=t,extraction_version='demo-v2') for i,t in enumerate(pages)],answers=answers,initial=initial))
def answer(question,label,text,quotes):return dict(question=question,label=label,answer=text,quotes=quotes)
def word(id,name,title,pages,answers,initial=True):
 d=Document();sec=d.sections[0];sec.page_width=Inches(7);sec.page_height=Inches(9);sec.top_margin=sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.65)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(12)
 for h in ['Title','Heading 1','Heading 2']:d.styles[h].font.color.rgb=RGBColor.from_string('262722')
 texts=[]
 for i,(heading,paras) in enumerate(pages):
  if i:d.add_page_break()
  p=d.add_paragraph('ACME  /  RESEARCH & STRATEGY');p.runs[0].font.size=Pt(9);p.runs[0].font.color.rgb=RGBColor.from_string('DF543F')
  d.add_heading(heading,0)
  d.add_paragraph('FY 2025  •  FICTIONAL DEMO DOCUMENT',style='Subtitle')
  for para in paras:d.add_paragraph(para)
  if i==0:
   t=d.add_table(rows=1, cols=2);t.style='Light Shading Accent 1';t.cell(0,0).text='SOURCE';t.cell(0,1).text='REVIEW STATUS';r=t.add_row();r.cells[0].text=title;r.cells[1].text='Prepared sample'
  texts.append('\n\n'.join(['ACME / RESEARCH & STRATEGY',heading,*paras]))
 sec.footer.paragraphs[0].text='ACME • Fictional sample for the Citely product demo'
 d.save(out/name);record(id,name,texts,answers,initial)
growth='Revenue increased by 24%, driven primarily by expansion within enterprise accounts.'
share='Enterprise customers represented 62% of total revenue.'
retention='Net revenue retention reached 118%, reflecting deeper adoption and sustained customer engagement.'
risks='The principal planning risks are slower enterprise purchase cycles and increased infrastructure costs.'
word('annual','Annual report 2025.docx','Annual review',[
 ('A year of meaningful growth.',[growth,share,retention,'The finance team paired commercial momentum with a disciplined approach to investment. This report brings the key findings into one place.']),
 ('A clear plan for the year ahead.',[risks,'The team will prioritize customer retention and platform reliability.','Investments will be introduced in stages, with quarterly checkpoints to review progress.','The next planning review is scheduled for 15 November 2025.'])],[
 answer('What drove revenue growth this year?','Revenue growth','Revenue grew 24%, driven by enterprise account expansion. Enterprise customers contributed 62% of total revenue.',[[1,growth],[1,share]]),
 answer('What are the main risks?','Main risks','The report identifies slower enterprise purchase cycles and increased infrastructure costs as the main planning risks.',[[2,risks]]),
 answer('What does the report say about retention?','Retention','Net revenue retention reached 118%, reflecting deeper adoption and sustained customer engagement.',[[1,retention]])])
w=Workbook();s=w.active;s.title='Summary';rows=[['ACME CASH FLOW','FY 2024','FY 2025'],['Operating cash flow',6.1,8.4],['Closing cash balance',10,14.2],['Investing outflow',1.6,2.1],['All amounts in USD millions; fictional sample','','']]
for row in rows:s.append(row)
s.column_dimensions['A'].width=48;s.column_dimensions['B'].width=18;s.column_dimensions['C'].width=18
for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='272822')
for row in s.iter_rows(min_row=2,max_row=4,min_col=2):
 for c in row:c.number_format='0.0';c.fill=PatternFill('solid',fgColor='FFF0D8')
s.freeze_panes='B2';a=w.create_sheet('Assumptions');a.append(['ASSUMPTION','DETAIL']);a.append(['Borrowing','No new borrowing is assumed next quarter.']);a.append(['Collections','Collections improved during the year.']);a.column_dimensions['A'].width=24;a.column_dimensions['B'].width=60
w.save(out/'Cash flow model.xlsx')
record('cash','Cash flow model.xlsx',['\n'.join('\t'.join(str(v) for v in r) for r in rows),'ASSUMPTION\tDETAIL\nBorrowing\tNo new borrowing is assumed next quarter.\nCollections\tCollections improved during the year.'],[
 answer('How strong is our cash position?','Cash position','Closing cash was $14.2 million, up from $10.0 million in FY 2024. Operating cash flow increased to $8.4 million.',[[1,'10\t14.2',dict(format='xlsx',sheet_name='Summary',sheet_index=0,range='B3:C3')],[1,'6.1\t8.4',dict(format='xlsx',sheet_name='Summary',sheet_index=0,range='B2:C2')]]),
 answer('What assumptions does the cash model use?','Assumptions','The model assumes no new borrowing next quarter and improved collections.',[[2,'No new borrowing is assumed next quarter.',dict(format='xlsx',sheet_name='Assumptions',sheet_index=1,range='B2')]])])
p=Presentation();p.slide_width=PI(10);p.slide_height=PI(5.625);pages=[];quotes=[]
for idx,(title,body) in enumerate([('Evidence into decisions','The board recommends staged investment while monitoring infrastructure costs.'),('Three priorities for the next quarter','Prioritize customer retention, platform reliability, and carefully sequenced product expansion.')]):
 slide=p.slides.add_slide(p.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=PC.from_string('23241F')
 for x,y,ww,hh,txt,size,color in [(.6,.4,8.7,.4,'ACME / BOARD UPDATE / FICTIONAL SAMPLE',12,'FF6347'),(.6,1.1,8.7,1.3,title,36,'F5F3E9'),(.6,2.9,8.7,1.5,body,22,'F5F3E9')]:
  shape=slide.shapes.add_textbox(PI(x),PI(y),PI(ww),PI(hh));shape.text_frame.word_wrap=True;run=shape.text_frame.paragraphs[0].add_run();run.text=txt;run.font.size=PP(size);run.font.name='Arial';run.font.color.rgb=PC.from_string(color)
 pages.append(title+'\n\n'+body);quotes.append([idx+1,body,dict(format='pptx',slide_index=idx,shape_id=shape.shape_id,quote=body,offset=0)])
p.save(out/'Board update.pptx')
record('board','Board update.pptx',pages,[answer('What does the board recommend?','Board decision','The board recommends staged investment while monitoring infrastructure costs.',[quotes[0]]),answer('What are the priorities?','Priorities','The priorities are customer retention, platform reliability and carefully sequenced product expansion.',[quotes[1]])])
memo='The research team interviewed twelve customers. Their top request was a clearer connection between answers and source passages.'
word('memo','Customer research memo.docx','Customer research',[('Listen before expanding.',[memo,'Participants wanted to keep useful passages beside their working notes.','These findings are fictional and included to demonstrate the upload workflow.'])],[answer('What did customers ask for?','Customer findings','The twelve sample customers asked for a clearer connection between answers and source passages.',[[1,memo]])],False)
for d in data:
 d['size']=(out/d['filename']).stat().st_size
 for i,a in enumerate(d['answers']):
  a['sources']=[]
  for j,q in enumerate(a.pop('quotes')):
   page,quote=q[:2];text=d['pages'][page-1]['text'];start=text.index(quote)
   a['sources'].append(dict(citation_id=f"{d['id']}-{i}-{j}",document_id=d['id'],filename=d['filename'],page=page,page_char_start=start,page_char_end=start+len(quote),quote=quote,content=quote,extraction_version='demo-v2',offset_unit='utf16',evidence_role='cited',**({'anchor':q[2]} if len(q)>2 else {})))
(root/'lib/studio-demo-data.json').write_text(json.dumps(data,indent=2))
print('Generated four fictional original files and matching citation fixtures.')
