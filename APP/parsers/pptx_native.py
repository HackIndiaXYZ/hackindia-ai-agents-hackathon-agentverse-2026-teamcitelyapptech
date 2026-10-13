"""Native .pptx extraction: chunks built from slide shapes, each knowing its shape.

Azure Document Intelligence returns a deck as flat text with no slide or shape
identity, so a citation can never point back at the box it came from. This
pass reads the deck with python-pptx and emits chunks that are already final:

- one "page" per slide (page = slide position + 1);
- one chunk per text-bearing shape, and one per table row, each prefixed with
  its slide's title ("Slide 2 — Quarter highlights") so a short bullet list
  still carries its context;
- `native_anchor` = {format: "pptx", slide_index, shape_id[, row]}, where
  shape_id is the shape's cNvPr id — the same id the browser renderer's text
  index reports — and `anchor_spans` locate the shape's own text (or each table
  cell's value) inside the chunk, so a cited quote can be narrowed to exact
  characters of one shape or cell.

Speaker notes, pictures, charts and SmartArt are not extracted: the viewer
cannot show a citation in them.
"""

from __future__ import annotations

import io
import os
import re

from langchain_core.documents import Document

from APP.parsers.base import DocumentParser

CHUNK_CHARS = 1000
MAX_SLIDES = int(os.getenv("PPTX_NATIVE_MAX_SLIDES", "500"))


def native_chunks_enabled() -> bool:
    # Read at call time: the parser singletons are built before .env is loaded.
    return os.getenv("NATIVE_CHUNKS_PPTX", "false").strip().lower() in ("1", "true", "yes")


class DeckTooLarge(ValueError):
    pass


def _clean(text: str) -> str:
    # python-pptx reports a line break inside a paragraph as a vertical tab.
    return (text or "").replace("\v", "\n").strip()


def _flat_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _flat_shapes(shape.shapes)       # its children carry the text and ids
        else:
            yield shape


def _reading_order(shape):
    return (shape.top or 0, shape.left or 0)


def _title_shape(slide):
    try:
        title = slide.shapes.title
    except (KeyError, AttributeError):
        return None
    return title if title is not None and _clean(title.text_frame.text) else None


def _header_labels(rows: list[list[str]]) -> list[str] | None:
    if len(rows) < 2:
        return None
    first = rows[0]
    filled = [t for t in first if t]
    if len(filled) >= 2 and all(not re.fullmatch(r"[-+]?[\d.,%€$£¥₹ ]+", t) for t in filled):
        return first
    return None


def _split_text(text: str, budget: int):
    """A long text box as pieces of at most `budget` characters, cut between
    paragraphs where possible. Yields (offset_in_text, piece)."""
    start = 0
    while start < len(text):
        end = min(len(text), start + budget)
        if end < len(text):
            cut = text.rfind("\n", start, end)
            if cut <= start:
                cut = text.rfind(" ", start, end)
            if cut > start:
                end = cut
        piece = text[start:end]
        if piece.strip():
            lead = len(piece) - len(piece.lstrip())
            yield start + lead, piece.strip()
        start = end
        while start < len(text) and text[start] in "\n ":
            start += 1


def parse_deck(raw_bytes: bytes, filename: str, chunk_chars: int = CHUNK_CHARS) -> list[Document]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(raw_bytes))
    if len(deck.slides) > MAX_SLIDES:
        raise DeckTooLarge(f"more than {MAX_SLIDES} slides")
    documents: list[Document] = []

    for slide_index, slide in enumerate(deck.slides):
        title = _title_shape(slide)
        title_text = _clean(title.text_frame.text) if title is not None else ""
        heading = f"Slide {slide_index + 1}" + (f" — {' '.join(title_text.split())}" if title_text else "") + "\n"
        units = []
        for shape in sorted(_flat_shapes(slide.shapes), key=_reading_order):
            if getattr(shape, "has_table", False) and shape.has_table:
                units.append(("table", shape))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and _clean(shape.text_frame.text):
                units.append(("text", shape))
        # The title already heads every chunk; it is a unit of its own only
        # when the slide has nothing else.
        if len(units) > 1 and title is not None:
            units = [u for u in units if u[1].shape_id != title.shape_id]

        def emit(text: str, spans: list, anchor: dict) -> None:
            documents.append(Document(page_content=text, metadata={
                "source": filename,
                "page": slide_index + 1,
                "slide_index": slide_index,
                "native_chunk": True,
                "native_anchor": {"format": "pptx", "slide_index": slide_index, **anchor},
                "anchor_spans": spans,
            }))

        for kind, shape in units:
            shape_id = int(shape.shape_id)
            if kind == "text":
                body = _clean(shape.text_frame.text)
                for offset, piece in _split_text(body, max(200, chunk_chars - len(heading))):
                    text = heading + piece
                    emit(text, [[len(heading), len(text), {"node_offset": offset}]], {"shape_id": shape_id})
                continue
            rows = [[_clean(cell.text) for cell in row.cells] for row in shape.table.rows]
            labels = _header_labels(rows)
            for r, values in enumerate(rows):
                if not any(values):
                    continue
                is_header = labels is not None and r == 0
                line, spans = f"Row {r + 1}{' (headers)' if is_header else ''}: ", []
                for c, value in enumerate(values):
                    if not value:
                        continue
                    if spans:
                        line += " | "
                    if labels is not None and not is_header and c < len(labels) and labels[c]:
                        line += f"{' '.join(labels[c].split())}: "
                    start = len(heading) + len(line)
                    line += value
                    spans.append([start, len(heading) + len(line), {"row": r, "col": c}])
                emit(heading + line, spans, {"shape_id": shape_id, "row": r})
    return documents


class LocalPptxParser(DocumentParser):
    """python-pptx, in process. Free, no network; enabled by NATIVE_CHUNKS_PPTX."""

    name = "local-pptx"
    extensions = (".pptx",)
    structured = True

    def is_available(self) -> bool:
        return native_chunks_enabled()

    def parse(self, raw_bytes: bytes, filename: str) -> list[Document]:
        return parse_deck(raw_bytes, filename)


def deck_units(raw_bytes: bytes) -> list[list]:
    """Every text-bearing shape and table cell as [slide_index, shape_id, row,
    col, text] — the identities citations use. Shared with the viewer's test,
    which checks the browser renderer reports the same ones."""
    from pptx import Presentation

    units = []
    for slide_index, slide in enumerate(Presentation(io.BytesIO(raw_bytes)).slides):
        for shape in sorted(_flat_shapes(slide.shapes), key=_reading_order):
            if getattr(shape, "has_table", False) and shape.has_table:
                for r, row in enumerate(shape.table.rows):
                    for c, cell in enumerate(row.cells):
                        if _clean(cell.text):
                            units.append([slide_index, int(shape.shape_id), r, c, _clean(cell.text)])
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and _clean(shape.text_frame.text):
                units.append([slide_index, int(shape.shape_id), None, None, _clean(shape.text_frame.text)])
    return units
