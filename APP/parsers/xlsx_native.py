"""Native .xlsx extraction: chunks built from cells, each knowing its cells.

Azure Document Intelligence returns a workbook as one flat page — sheet names
become ordinary text lines, dates become serial numbers (2026-03-31 → 46112)
and wide sheets are read with their columns interleaved — so a citation can
never point back at a cell. This pass reads the workbook itself with openpyxl
and emits chunks that are already final:

- one "page" per sheet (page = sheet position + 1), so the extracted-text view
  and page numbers keep working;
- each chunk is a run of whole rows from one sheet, rendered as
  "Row 5: Metric: Gross margin | Actual: 41.2%" so every value keeps its column
  header — the association a table needs to be answerable;
- `native_anchor` = {format, sheet_name, sheet_index, range}, and
  `anchor_spans` = [[start, end, {"range": A1}], ...] locating every cell inside
  the chunk text, so a cited quote narrows to exactly the cells it covers.

Identifiers match what SheetJS reports in the browser (same sheet order, same
A1 addresses, merges anchored on their top-left cell), so the viewer can
highlight a citation without re-deriving anything.

A formula cell keeps its saved value as the cell's own anchored text (what the
browser viewer shows, so both sides still address the same units), and its
expression is recorded beside it, outside the anchor: "41.2% [formula =B4*2]".
A formula with no saved value is recorded as "[formula =B4*2; no saved value]"
and is not an anchored unit, since the viewer shows nothing in that cell. No
formula is recalculated: a saved value is whatever the spreadsheet program
stored, and its freshness is not verified (chunks say so in their metadata).
"""

from __future__ import annotations

import datetime as dt
import io
import os
import re
import zipfile

from langchain_core.documents import Document

from APP.parsers.base import DocumentParser

# Guards against pathological workbooks before openpyxl loads everything into
# memory. Above either limit the router falls through to the next tier.
MAX_SHEET_XML_BYTES = int(os.getenv("XLSX_NATIVE_MAX_XML_BYTES", str(200 * 1024 * 1024)))
MAX_CELLS = int(os.getenv("XLSX_NATIVE_MAX_CELLS", "300000"))
CHUNK_CHARS = 1000


def native_chunks_enabled() -> bool:
    # Read at call time: the parser singletons are built before .env is loaded.
    return os.getenv("NATIVE_CHUNKS_XLSX", "false").strip().lower() in ("1", "true", "yes")


class WorkbookTooLarge(ValueError):
    pass


# ---------------------------------------------------------------------------
# Cell values as text
# ---------------------------------------------------------------------------

_CURRENCY = "$€£¥₹"


def _currency_of(section: str) -> str:
    bracket = re.search(r"\[\$([^\-\]]+)", section)
    if bracket:
        return bracket.group(1)
    for literal in re.findall(r'"([^"]*)"', section):
        for symbol in _CURRENCY:
            if symbol in literal:
                return symbol
    bare = re.sub(r'"[^"]*"|\[[^\]]*\]', "", section)
    return next((s for s in _CURRENCY if s in bare), "")


def _general(value: float) -> str:
    if float(value).is_integer() and abs(value) < 1e15:
        return str(int(value))
    return f"{value:.10g}"


def format_number(value: float, number_format: str | None) -> str:
    """The number roughly as Excel displays it: percent, decimals, thousands
    separators and a currency symbol. Anything more exotic falls back to
    Excel's "General" rendering, which never misstates the value."""
    fmt = number_format or "General"
    if fmt == "General":
        return _general(value)
    section = fmt.split(";")[0]
    core = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.', "", section)
    if not re.search(r"[0#?]", core) or re.search(r"[Ee][+-]|/", core):
        return _general(value)
    percent = "%" in core
    if percent:
        value = value * 100
    decimals_match = re.search(r"\.([0#?]+)", core)
    decimals = len(decimals_match.group(1)) if decimals_match else 0
    thousands = "," in core.split(".")[0]
    text = f"{value:,.{decimals}f}" if thousands else f"{value:.{decimals}f}"
    sign = ""
    if text.startswith("-"):
        sign, text = "-", text[1:]
    return f"{sign}{_currency_of(section)}{text}{'%' if percent else ''}"


def cell_text(value, number_format: str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds" if value.second else "minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat(timespec="seconds" if value.second else "minutes")
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, (int, float)):
        return format_number(value, number_format)
    return re.sub(r"\s+", " ", str(value)).strip()


# ---------------------------------------------------------------------------
# Workbook → chunks
# ---------------------------------------------------------------------------

def _check_size(raw_bytes: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
        total = sum(i.file_size for i in archive.infolist() if i.filename.startswith("xl/worksheets/"))
    if total > MAX_SHEET_XML_BYTES:
        raise WorkbookTooLarge(f"worksheets expand to {total} bytes")


def _header_row(rows: list[tuple[int, list[tuple[int, str, str, str | None]]]]) -> int | None:
    """The first row that looks like column headings: at least two filled
    cells, none of them numeric, with data below it."""
    for position, (row_number, cells) in enumerate(rows[:-1]):
        valued = [text for _, text, _, _ in cells if text]
        if len(valued) >= 2 and all(not re.fullmatch(r"[-+]?[\d.,%€$£¥₹ ]+", text) for text in valued):
            return row_number
    return None


def bounding_range(ranges: list[str]) -> str:
    """Smallest A1 range covering every given cell or merge range."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    boxes = [range_boundaries(r) for r in ranges]
    min_col = min(b[0] for b in boxes)
    min_row = min(b[1] for b in boxes)
    max_col = max(b[2] for b in boxes)
    max_row = max(b[3] for b in boxes)
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _sheet_rows(ws, formulas=None) -> list[tuple[int, list[tuple[int, str, str, str | None]]]]:
    """Non-empty rows as (row_number, [(column, text, a1_or_merge_range, formula_note)]).
    text is the saved value (empty for a formula with none); formula_note is
    None for an ordinary cell."""
    from openpyxl.utils.cell import get_column_letter

    merge_at: dict[tuple[int, int], str] = {}
    covered: set[tuple[int, int]] = set()
    for merged in ws.merged_cells.ranges:
        merge_at[(merged.min_row, merged.min_col)] = merged.coord
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                if (r, c) != (merged.min_row, merged.min_col):
                    covered.add((r, c))

    rows = []
    for row in ws.iter_rows():
        cells = []
        for cell in row:
            # openpyxl already blanks merge-covered cells on load; checked
            # anyway so a covered value can never become a second unit.
            formula = formulas.cell(cell.row, cell.column) if formulas is not None else None
            is_formula = formula is not None and formula.data_type == "f"
            if (cell.value is None and not is_formula) or (cell.row, cell.column) in covered:
                continue
            text = cell_text(cell.value, cell.number_format) if cell.value is not None else ""
            note = None
            if is_formula:
                expression = getattr(formula.value, "text", formula.value)
                note = f"formula {expression}" if text else f"formula {expression}; no saved value"
            if not text and note is None:
                continue
            address = merge_at.get((cell.row, cell.column)) or f"{get_column_letter(cell.column)}{cell.row}"
            cells.append((cell.column, text, address, note))
        if cells:
            rows.append((row[0].row, cells))
    return rows


def _row_lines(row_number: int, cells, labels: dict[int, str], is_header: bool, budget: int):
    """One row as one or more lines, each with its cell spans (relative to the
    line). A row too long for a chunk is split between cells."""
    prefix = f"Row {row_number}{' (headers)' if is_header else ''}: "
    line, spans = prefix, []
    for column, text, address, note in cells:
        lead = "" if is_header or column not in labels else f"{labels[column]}: "
        piece = f"{lead}{text}" if text else f"{lead}[{note}]"
        tail = f" [{note}]" if note and text else ""
        separator = "" if line == prefix else " | "
        if spans and len(line) + len(separator) + len(piece) + len(tail) > budget:
            yield line, spans
            line, spans, separator = prefix, [], ""
        start = len(line) + len(separator)
        line = f"{line}{separator}{piece}"
        if text:
            spans.append([start, len(line), {"range": address}])
        line += tail          # the formula is recorded beside the value, outside its anchor
    if spans or line != prefix:
        yield line, spans


def parse_workbook(raw_bytes: bytes, filename: str, chunk_chars: int = CHUNK_CHARS) -> list[Document]:
    import openpyxl
    from openpyxl.worksheet.worksheet import Worksheet

    _check_size(raw_bytes)
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=False)
    formula_wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False, read_only=False)
    documents: list[Document] = []
    seen_cells = 0
    try:
        for sheet_index, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            if not isinstance(ws, Worksheet):  # chartsheets carry no cells
                continue
            rows = _sheet_rows(ws, formula_wb[sheet_name])
            seen_cells += sum(len(cells) for _, cells in rows)
            if seen_cells > MAX_CELLS:
                raise WorkbookTooLarge(f"more than {MAX_CELLS} filled cells")
            if not rows:
                continue
            header = _header_row(rows)
            labels = {}
            if header is not None:
                labels = {column: text for column, text, _, _ in dict(rows)[header] if text}

            heading = f"Sheet: {sheet_name}\n"
            text, spans = heading, []

            def flush():
                nonlocal text, spans
                if spans:
                    documents.append(Document(page_content=text.rstrip("\n"), metadata={
                        "source": filename,
                        "page": sheet_index + 1,
                        "sheet_name": sheet_name,
                        "sheet_index": sheet_index,
                        "native_chunk": True,
                        "native_anchor": {"format": "xlsx", "sheet_name": sheet_name, "sheet_index": sheet_index,
                                          "range": bounding_range([s[2]["range"] for s in spans])},
                        "anchor_spans": spans,
                        **({"formula_values": "saved_not_recalculated"} if " [formula " in text or "[formula " in text else {}),
                    }))
                text, spans = heading, []

            for row_number, cells in rows:
                # Rows above the header (a title, a note) are not part of the table.
                row_labels = labels if header is not None and row_number > header else {}
                for line, line_spans in _row_lines(row_number, cells, row_labels, row_number == header,
                                                   chunk_chars - len(heading)):
                    if spans and len(text) + len(line) + 1 > chunk_chars:
                        flush()
                    offset = len(text)
                    text += line + "\n"
                    spans.extend([offset + s, offset + e, ref] for s, e, ref in line_spans)
            flush()
    finally:
        wb.close()
        formula_wb.close()
    return documents


class LocalXlsxParser(DocumentParser):
    """openpyxl, in process. Free, no network; enabled by NATIVE_CHUNKS_XLSX."""

    name = "local-xlsx"
    extensions = (".xlsx",)
    # Output is structured by construction; the OCR-oriented density check
    # would wrongly reject a small but perfectly extracted sheet.
    structured = True

    def is_available(self) -> bool:
        return native_chunks_enabled()

    def parse(self, raw_bytes: bytes, filename: str) -> list[Document]:
        return parse_workbook(raw_bytes, filename)
