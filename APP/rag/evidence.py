"""Map model-selected evidence to exact canonical page ranges.

Extraction/index offsets use Python code points. The API explicitly sends
UTF-16 code units, the indexing convention used by browser text.slice().
"""
import json
import re
from uuid import NAMESPACE_URL, uuid5

from APP.rag.chunking import _normalise_with_map
from APP.schemas import SourceChunk

LIMITED_CONTEXT = "I couldn't find enough information in the retrieved passages to answer that. This does not establish that the information is absent from the document."
UNVERIFIED_ANSWER = "I couldn't verify the generated answer against the source passages. Please try rephrasing your question."


def utf16_offset(text: str, index: int) -> int:
    return len(text[:index].encode('utf-16-le')) // 2


def codepoint_offset(text: str, index: int) -> int | None:
    count = 0
    for i, char in enumerate(text):
        if count == index:
            return i
        count += 2 if ord(char) > 0xFFFF else 1
    return len(text) if count == index else None


def unique_quote_range(text: str, quote: str) -> tuple[int, int] | None:
    normalized, offsets = _normalise_with_map(text)
    needle = _normalise_with_map(quote)[0].strip()
    if not needle:
        return None
    start = normalized.find(needle)
    if start < 0 or normalized.find(needle, start + 1) >= 0:
        return None
    return offsets[start], offsets[start + len(needle) - 1] + 1


def _bounding_range(ranges: list[str]) -> str:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    boxes = [range_boundaries(r) for r in ranges]
    start = f"{get_column_letter(min(b[0] for b in boxes))}{min(b[1] for b in boxes)}"
    end = f"{get_column_letter(max(b[2] for b in boxes))}{max(b[3] for b in boxes)}"
    return start if start == end else f"{start}:{end}"


def native_anchor_for(meta: dict, text: str, quote: str | None) -> dict | None:
    """The chunk's native anchor, narrowed to the units the quote covers.

    A spreadsheet chunk spans several rows; a quote naming one value should
    light up that cell, not the whole block. Narrowing only happens when the
    quote is found exactly once in the chunk; otherwise the whole chunk's
    anchor is returned, which is still correct, just less precise.
    """
    anchor = meta.get("native_anchor")
    if not anchor:
        return None
    anchor = dict(anchor)
    spans = meta.get("anchor_spans") or []
    found = unique_quote_range(text, quote) if quote and spans else None
    if not found:
        return anchor
    start, end = found
    hits = [(s, e, ref) for s, e, ref in spans if s < end and e > start]
    if not hits:
        return anchor
    if anchor.get("format") == "xlsx":
        cells = list(dict.fromkeys(ref["range"] for _, _, ref in hits))
        anchor["range"] = _bounding_range(cells)
        anchor["cells"] = cells
    elif len(hits) == 1:
        # One shape or one table cell: keep exactly the cited characters of
        # it, and where they start in that shape's own text.
        s, e, ref = hits[0]
        a, b = max(s, start), min(e, end)
        if "col" in ref:
            anchor["col"] = ref["col"]
        anchor["quote"] = text[a:b]
        anchor["offset"] = ref.get("node_offset", 0) + (a - s)
    else:
        anchor["cells"] = [[ref["row"], ref["col"]] for _, _, ref in hits if "col" in ref]
    return anchor


def _reply_payload(raw: str):
    """The model's reply as a dict. JSON, as asked; but the model sometimes
    answers in YAML-style text ("answer_status: supported / answer: | ..."),
    which used to be shown to the reader verbatim, header and all. Raises
    ValueError when the reply is neither."""
    text = re.sub(r'^```(?:json|yaml)?\s*|\s*```$', '', raw.strip())
    if not re.match(r'answer_status\s*:', text):
        return json.loads(text)
    import textwrap
    import yaml
    try:
        loaded = yaml.safe_load(text)
    except yaml.YAMLError:
        loaded = None
    if isinstance(loaded, dict) and isinstance(loaded.get('answer'), str):
        return loaded
    # A quote with a colon in it can break the YAML: take the answer text between
    # its header and any citations block, and the citations that can be read.
    body = re.search(r'\n\s*answer\s*:\s*[|>]?[-+]?[ \t]*\n?(.*?)(?:\n\s*citations\s*:|\Z)', text, re.S)
    if not body:
        raise ValueError('unreadable reply')
    status = re.match(r'answer_status\s*:\s*["\']?(\w+)', text)
    citations = [{'id': m[1], 'quote': m[2].strip().strip('"\'')}
                 for m in re.finditer(r'id\s*:\s*["\']?(C\d+)["\']?\s*\n\s*quote\s*:\s*(.+)', text)]
    return {'answer_status': status[1] if status else 'supported', 'answer': textwrap.dedent(body[1]).strip(), 'citations': citations}


def selected_evidence(raw: str, kept: list) -> tuple[str, list[tuple[int, str | None]]]:
    """Never mistake the context inventory for evidence cited by the answer."""
    quotes = {}
    references = {}
    # The model sometimes puts its status on a line of its own ahead of ordinary prose:
    # "supported\n\n**Answer:** ...". That word is bookkeeping, never part of the answer.
    status_line = re.match(r'\s*(supported|insufficient_evidence)[ \t]*\n', raw)
    if status_line:
        if status_line[1] == 'insufficient_evidence':
            return LIMITED_CONTEXT, []
        raw = raw[status_line.end():].lstrip()
    answer = raw
    try:
        payload = _reply_payload(raw)
        if isinstance(payload, dict) and isinstance(payload.get('answer'), str):
            if payload.get('answer_status') == 'insufficient_evidence':
                return LIMITED_CONTEXT, []
            answer = payload['answer']
            for citation in payload.get('citations', []) or []:
                if isinstance(citation, dict) and isinstance(citation.get('quote'), str) and citation['quote'].strip():
                    ident = str(citation.get('id', ''))
                    if re.fullmatch(r'E[1-9]\d*', ident):
                        passage = re.fullmatch(r'C([1-9]\d*)', str(citation.get('passage_id', '')))
                        if ident in references or not passage or int(passage[1]) > len(kept):
                            return UNVERIFIED_ANSWER, []
                        references[ident] = (int(passage[1]) - 1, citation['quote'])
                        continue
                    quotes.setdefault(ident, [])
                    if citation['quote'] not in quotes[ident]:
                        quotes[ident].append(citation['quote'])
            # Some providers explicitly select evidence in the JSON array
            # but omit inline markers. Attach those selected IDs, never the
            # retrieval inventory, as answer-level references.
            if not re.search(r'\[(?:C|E)\d+\]', answer):
                markers = [ident for ident in quotes if re.fullmatch(r'C[1-9]\d*', ident)
                           and int(ident[1:]) <= len(kept)]
                if markers:
                    answer += ' ' + ''.join(f'[{ident}]' for ident in markers)
    except (ValueError, TypeError):
        if raw.lstrip().startswith(('{', '```')):
            # A reply that looks like JSON but can't be read (usually cut off by the
            # output limit) says nothing about the document. It is not "the document
            # lacks this": it has no usable answer, so it is retried, and if that fails
            # too the reader is told the answer couldn't be verified.
            return '', []
    chosen = []
    # Evidence IDs identify individual quotes, independently of passage IDs.
    reference_numbers = {}
    def evidence_marker(match):
        ident = match[1]
        if ident not in reference_numbers:
            chosen.append(references[ident])
            reference_numbers[ident] = len(chosen)
        return f"[E{reference_numbers[ident]}]"
    if any(ident not in references for ident in re.findall(r'\[(E\d+)\]', answer)):
        return UNVERIFIED_ANSWER, []
    answer = re.sub(r'\[(E\d+)\]', evidence_marker, answer)
    for marker in dict.fromkeys(re.findall(r'\[C(\d+)\]', answer)):
        index = int(marker) - 1
        if 0 <= index < len(kept):
            chosen.extend((index, quote) for quote in quotes.get('C' + marker, [None]))
    # Compatibility with older providers: filename/page is only usable when
    # it identifies ONE supplied passage. Never turn one page marker into ten.
    if not chosen:
        for name, page in re.findall(r'\[Source:\s*(.*?)\s*\|\s*Page:\s*(.*?)\]', answer):
            matches = [i for i, (doc, _) in enumerate(kept)
                       if doc.metadata.get('source', 'unknown') == name and str(doc.metadata.get('page', '?')) == page]
            if len(matches) == 1 and (matches[0], None) not in chosen:
                chosen.append((matches[0], None))
    return answer, chosen


def source_for(document, score, page: dict | None, quote: str | None = None) -> SourceChunk:
    meta = document.metadata
    source = SourceChunk(chunk_id=meta.get('chunk_id'), document_id=meta.get('document_id'),
                         source=meta.get('source', 'unknown'), page=meta.get('page'), score=score,
                         content=document.page_content, quote=quote or document.page_content,
                         extraction_version=meta.get('extraction_version'), offset_unit='utf16', evidence_role='cited')
    source.region_id = meta.get('region_id')
    source.evidence_record_id = meta.get('evidence_record_id')
    source.provenance = meta.get('provenance', 'native')
    source.content_type = meta.get('content_type', 'text')
    source.evidence_location = meta.get('evidence_location')
    if source.evidence_record_id and source.provenance != 'native':
        location = source.evidence_location or {}
        if location.get('format') in ('pptx', 'pdf'):
            source.anchor = dict(location, evidence_region=True)
        source.citation_id = str(uuid5(NAMESPACE_URL, json.dumps([source.evidence_record_id, source.quote])))
        return source
    preview = bool(page and page.get('extraction_version', '').startswith('legacy-preview-'))
    if page and page.get('extraction_version') and (preview or meta.get('extraction_version') in (None, '', page['extraction_version'])):
        text = page['text']
        start, end = meta.get('page_char_start'), meta.get('page_char_end')
        valid = not preview and type(start) is int and type(end) is int and 0 <= start < end <= len(text)
        if valid and _normalise_with_map(text[start:end])[0].strip() != _normalise_with_map(document.page_content)[0].strip():
            valid = False
        if not valid:
            found = unique_quote_range(text, document.page_content)
            start, end = found if found else (None, None)
        if start is not None:
            if quote:
                found = unique_quote_range(text[start:end], quote)
                if found:
                    start, end = start + found[0], start + found[1]
                else:
                    start = end = None
            if start is not None:
                source.page_char_start = utf16_offset(text, start)
                source.page_char_end = utf16_offset(text, end)
                source.quote = text[start:end]
                source.extraction_version = page['extraction_version']
    source.anchor = native_anchor_for(meta, document.page_content, quote)
    identity = [source.document_id, source.page, source.extraction_version,
                source.page_char_start, source.page_char_end, source.quote]
    source.citation_id = str(uuid5(NAMESPACE_URL, json.dumps(identity, ensure_ascii=False)))
    return source


def checked_quotes(selected: list, kept: list) -> tuple[tuple[str, ...], bool]:
    """Check evidence against supplied passages, independently of model status.

    This establishes quotation provenance, not semantic entailment. Canonical
    page/version/range validation still happens in source_for.
    """
    quotes = []
    valid = True
    for index, quote in selected:
        text = kept[index][0].page_content
        if quote is None:  # Older providers cite an entire identified passage.
            quote = text
        needle = _normalise_with_map(quote)[0].strip()
        if not needle or needle not in _normalise_with_map(text)[0]:
            valid = False
            continue
        quotes.append(quote)
    return tuple(quotes), valid


def near_quote(text: str, quote: str, share: float = 0.85) -> bool:
    """True when a quote is not exact but nearly is: at least `share` of its
    words appear in the passage in the same order (punctuation, case and a
    changed word or two aside), and it has at least four words. Exact quotes
    are False here: they are checked, and used, as they are."""
    words = lambda t: re.findall(r"[\w'’]+", _normalise_with_map(t)[0].casefold())
    q, p = words(quote), words(text)
    if len(q) < 4 or checked_quotes_exact(text, quote):
        return False
    # Longest common subsequence of words, bounded: quotes are short.
    q = q[:60]
    prev = [0] * (len(q) + 1)
    for w in p:
        cur = [0]
        for j, qw in enumerate(q, 1):
            cur.append(prev[j - 1] + 1 if w == qw else max(prev[j], cur[j - 1]))
        prev = cur
    return prev[-1] >= share * len(q)


def checked_quotes_exact(text: str, quote: str) -> bool:
    needle = _normalise_with_map(quote)[0].strip()
    return bool(needle) and needle in _normalise_with_map(text)[0]


def verified_excerpt_answer(selected: list, kept: list) -> tuple[str, list]:
    """After one failed repair, return only verified, selected source excerpts.

    Do not retain generated claims when their supporting quotations failed.
    This is deliberately distinct from a semantic verification of those claims.
    """
    verified, excerpts = [], []
    for index, quote in selected:
        quotes, valid = checked_quotes([(index, quote)], kept)
        if valid and quotes and (index, quotes[0]) not in verified:
            verified.append((index, quotes[0]))
            excerpts.append(f'- {quotes[0]} [E{len(verified)}]')
    if not verified:
        return UNVERIFIED_ANSWER, []
    return 'Supporting passages from your document:\n\n' + '\n\n'.join(excerpts), verified


def bounded_absence(answer: str, source_quotes: tuple[str, ...] = ()) -> str:
    # Restrict this check to explicit claims about a document's coverage.
    # Facts like "the product does not include PII detection" are ordinary
    # supported negative answers, not claims that retrieval searched all pages.
    scope = r"(?:the|this|your|entire|whole)\s+(?:document|file|report|synopsis|pdf)"
    denial = (rf"\b{scope}\s+(?:does not|doesn't)\s+(?:explicitly\s+)?(?:mention|contain|discuss|describe|include)\b"
              rf"|\b(?:no (?:mention|reference|information|evidence)|(?:is not|isn't) (?:mentioned|found|present))\b[^.!?\n]*\b(?:in|throughout|anywhere in)\s+{scope}\b")
    evidence = tuple(_normalise_with_map(q)[0].strip().casefold() for q in source_quotes)
    for sentence in re.split(r'(?<=[.!?])\s+|\n', answer):
        if re.search(denial, sentence, re.I):
            claim = re.sub(r'\[(?:C|E)\d+\]', '', sentence).strip().strip(' ."“”')
            normalized = _normalise_with_map(claim)[0].casefold()
            if not any(normalized in quote for quote in evidence):
                return LIMITED_CONTEXT
    return answer


# --------------------------------------------------------------------------- #
# Pinning a written summary to the sentences it came from
# --------------------------------------------------------------------------- #

_STOP = frozenset("""a an the and or but if then than of to in on at by for with from as is are was were be been being it its this that these those
which who whom what when where how why not no can could will would should may might must do does did done has have had having
into over under about across after before between during each per such their there they them our we you your his her also more most
other some any all both only very same so up out one two three""".split())


def _content_words(text: str) -> set[str]:
    words = re.findall(r"[a-z0-9][a-z0-9'’-]*", _normalise_with_map(text)[0].casefold())
    return {w[:-1] if len(w) > 4 and w.endswith("s") else w for w in words if len(w) > 2 and w not in _STOP}


def _passage_sentences(text: str) -> list[str]:
    """The passage's sentences (and lines), exactly as written, long enough to cite."""
    found = []
    for match in re.finditer(r"[^.!?\n]+(?:[.!?]+|$)", text, re.M):
        sentence = match.group().strip()
        if len(sentence.split()) >= 4:
            found.append(sentence)
    return found


def best_sentence(passage: str, claim: str) -> str | None:
    """The sentence of the passage a written claim most plainly draws on, or
    None when no sentence shares enough with it (or the best one appears more
    than once, so it couldn't be located). Plain word overlap, no model call."""
    claim_words = _content_words(claim)
    if len(claim_words) < 2:
        return None
    best, best_score = None, 0.0
    for sentence in _passage_sentences(passage):
        words = _content_words(sentence)
        shared = len(claim_words & words)
        if shared < 2 or not words:
            continue
        score = shared / len(words) + shared / len(claim_words)
        if score > best_score:
            best, best_score = sentence, score
    if best is None or best_score < 0.5 or unique_quote_range(passage, best) is None:
        return None
    return best


def pin_summary_citations(answer: str, kept: list) -> tuple[str, list[tuple[str, int, str | None]]]:
    """For a written summary: each passage marker is tied to the sentence of
    that passage closest to the text just before the marker, so a citation
    opens at that sentence. Markers get their own keys ([Cp1], [Cp2], ...)
    because the same passage can back different points with different
    sentences. Returns the re-keyed answer and (key, passage index, sentence
    or None for the whole passage)."""
    pins: list[tuple[str, int, str | None]] = []
    out, last = [], 0
    for group in re.finditer(r"(?:\s*\[C\d+\])+", answer):
        before = answer[last:group.start()]
        claim = re.split(r"\n|(?<=[.!?])\s+", before.strip())[-1] if before.strip() else ""
        keys = []
        for number in re.findall(r"\[C(\d+)\]", group.group()):
            index = int(number) - 1
            if not 0 <= index < len(kept):
                continue
            sentence = best_sentence(kept[index][0].page_content, claim)
            key = f"p{len(pins) + 1}"
            pins.append((key, index, sentence))
            keys.append(f"[C{key}]")
        out.append(before)
        out.append(re.match(r"\s*", group.group()).group() + "".join(keys) if keys else "")
        last = group.end()
    out.append(answer[last:])
    return "".join(out), pins
