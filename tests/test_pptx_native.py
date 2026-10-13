"""Native slide extraction: shapes, tables, anchors, routing, citations."""

import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from langchain_core.documents import Document
from pptx import Presentation
from pptx.util import Inches

import APP.rag  # noqa: F401 — import order the app uses (APP.parsers ↔ APP.rag)
from APP.parsers import pptx_native, router
from APP.parsers.pptx_native import DeckTooLarge, deck_units, parse_deck
from APP.rag.chunking import build_page_texts, finalize_native_chunks
from APP.rag.evidence import native_anchor_for

FIXTURES = Path(__file__).resolve().parents[1] / "Citely-Modernist-Next-Final" / "tests" / "fixtures" / "pptx"


def _deck(build) -> bytes:
    deck = Presentation()
    build(deck)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _spans(doc):
    return [(doc.page_content[s:e], ref) for s, e, ref in doc.metadata["anchor_spans"]]


class ParseDeckTest(unittest.TestCase):
    def test_shapes_become_chunks_headed_by_their_slide_title(self):
        def build(deck):
            s = deck.slides.add_slide(deck.slide_layouts[1])
            s.shapes.title.text = "Quarter highlights"
            s.placeholders[1].text_frame.text = "Churn fell to 3.1%."
            s.placeholders[1].text_frame.add_paragraph().text = "Safety first."
            only = deck.slides.add_slide(deck.slide_layouts[5])
            only.shapes.title.text = "Section break"
        docs = parse_deck(_deck(build), "d.pptx")
        self.assertEqual([d.page_content for d in docs], [
            "Slide 1 — Quarter highlights\nChurn fell to 3.1%.\nSafety first.",
            "Slide 2 — Section break\nSection break",           # a lone title is kept
        ])
        self.assertEqual(docs[0].metadata["native_anchor"], {"format": "pptx", "slide_index": 0, "shape_id": 3})
        self.assertEqual(_spans(docs[0]), [("Churn fell to 3.1%.\nSafety first.", {"node_offset": 0})])
        self.assertEqual([d.metadata["page"] for d in docs], [1, 2])

    def test_table_rows_keep_headers_and_spans_cover_each_cell_value(self):
        def build(deck):
            s = deck.slides.add_slide(deck.slide_layouts[5])
            s.shapes.title.text = "Fleet"
            table = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
            for r, row in enumerate([("Region", "Trucks"), ("Benelux", "140"), ("DACH", "171")]):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
        docs = parse_deck(_deck(build), "t.pptx")
        self.assertEqual([d.page_content.splitlines()[1] for d in docs],
                         ["Row 1 (headers): Region | Trucks", "Row 2: Region: Benelux | Trucks: 140", "Row 3: Region: DACH | Trucks: 171"])
        self.assertEqual(_spans(docs[2]), [("DACH", {"row": 2, "col": 0}), ("171", {"row": 2, "col": 1})])
        self.assertEqual(docs[2].metadata["native_anchor"]["row"], 2)

    def test_grouped_shapes_are_found_with_their_own_ids(self):
        def build(deck):
            s = deck.slides.add_slide(deck.slide_layouts[6])
            group = s.shapes.add_group_shape()
            box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            box.text_frame.text = "Inside a group"
        units = deck_units(_deck(build))
        self.assertEqual(len(units), 1)
        self.assertEqual(units[0][4], "Inside a group")
        self.assertGreater(units[0][1], 1)                       # the child's id, not the group's

    def test_long_text_splits_and_offsets_point_into_the_shape_text(self):
        body = "\n".join(f"Paragraph {i} " + "words " * 30 for i in range(12))
        def build(deck):
            s = deck.slides.add_slide(deck.slide_layouts[6])
            box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
            box.text_frame.text = body
        docs = parse_deck(_deck(build), "long.pptx", chunk_chars=500)
        self.assertGreater(len(docs), 4)
        node = deck_units(_deck(build))[0][4]
        for doc in docs:
            self.assertLessEqual(len(doc.page_content), 500)
            (piece, ref), = _spans(doc)
            self.assertEqual(node[ref["node_offset"]:ref["node_offset"] + len(piece)], piece)

    def test_line_breaks_inside_a_paragraph_read_as_newlines(self):
        def build(deck):
            s = deck.slides.add_slide(deck.slide_layouts[6])
            box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text_frame.paragraphs[0].add_run().text = "first"
            box.text_frame.paragraphs[0].add_line_break()
            box.text_frame.paragraphs[0].add_run().text = "second"
        self.assertEqual(deck_units(_deck(build))[0][4], "first\nsecond")

    def test_oversized_decks_are_refused_so_the_router_falls_back(self):
        def build(deck):
            for _ in range(4):
                deck.slides.add_slide(deck.slide_layouts[6])
        with mock.patch.object(pptx_native, "MAX_SLIDES", 3), self.assertRaises(DeckTooLarge):
            parse_deck(_deck(build), "big.pptx")

    def test_native_chunks_are_located_exactly_in_their_slide_text(self):
        docs = parse_deck((FIXTURES / "atlas_board_update.pptx").read_bytes(), "a.pptx")
        pages = build_page_texts(docs)
        for chunk in finalize_native_chunks(docs):
            start, end = chunk.metadata["page_char_start"], chunk.metadata["page_char_end"]
            self.assertEqual(pages[chunk.metadata["page"]][start:end], chunk.page_content)


class RoutingTest(unittest.TestCase):
    def setUp(self):
        self.small = _deck(lambda d: setattr(d.slides.add_slide(d.slide_layouts[5]).shapes.title, "text", "Hi"))
        for target, attr, value in [(router._azure_di, "is_available", True),
                                    (router._azure_di, "parse", [Document(page_content="x" * 500, metadata={"page": 1})])]:
            patcher = mock.patch.object(target, attr, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_flag_off_keeps_azure_document_intelligence(self):
        with mock.patch.dict(os.environ, {"NATIVE_CHUNKS_PPTX": "false"}):
            self.assertEqual(router.parse_document(self.small, "a.pptx").parser, "azure-document-intelligence")

    def test_flag_on_uses_native_and_a_failure_falls_back(self):
        with mock.patch.dict(os.environ, {"NATIVE_CHUNKS_PPTX": "true"}):
            self.assertEqual(router.parse_document(self.small, "a.pptx").parser, "local-pptx")
            with mock.patch.object(router._local_pptx, "parse", side_effect=DeckTooLarge("big")):
                result = router.parse_document(self.small, "a.pptx")
        self.assertEqual((result.parser, result.fallback_used), ("azure-document-intelligence", True))


class AnchorsTest(unittest.TestCase):
    def setUp(self):
        self.docs = parse_deck((FIXTURES / "atlas_board_update.pptx").read_bytes(), "a.pptx")
        self.by = {(d.metadata["native_anchor"]["slide_index"], d.metadata["native_anchor"]["shape_id"],
                    d.metadata["native_anchor"].get("row")): d for d in self.docs}

    def test_a_quote_narrows_to_its_characters_and_their_offset_in_the_shape(self):
        doc = self.by[(1, 3, None)]
        quote = "now handles 1,150 pallets a day"
        anchor = native_anchor_for(doc.metadata, doc.page_content, quote)
        self.assertEqual(anchor["quote"], quote)
        node = [u for u in deck_units((FIXTURES / "atlas_board_update.pptx").read_bytes()) if u[:2] == [1, 3]][0][4]
        self.assertEqual(node[anchor["offset"]:anchor["offset"] + len(quote)], quote)

    def test_a_quote_reaching_into_the_slide_heading_keeps_only_the_shape_part(self):
        doc = self.by[(3, 4, None)]
        anchor = native_anchor_for(doc.metadata, doc.page_content, "Principal risks\nDriver attrition")
        self.assertEqual((anchor["quote"], anchor["offset"]), ("Driver attrition", 0))

    def test_table_quotes_narrow_to_a_cell_or_list_the_cells_of_the_row(self):
        doc = self.by[(2, 3, 4)]
        one = native_anchor_for(doc.metadata, doc.page_content, "171")
        self.assertEqual((one["row"], one["col"], one["quote"], one["offset"]), (4, 1, "171", 0))
        row = native_anchor_for(doc.metadata, doc.page_content, "Region: DACH | Trucks: 171 | Utilisation: 90%")
        self.assertEqual(row["cells"], [[4, 0], [4, 1], [4, 2]])
        self.assertNotIn("quote", row)


class IdentityContractTest(unittest.TestCase):
    """The viewer must find the shapes the backend cites. This half pins what
    python-pptx emits; tests/browser/viewer-pptx.spec.ts checks the slide
    renderer's text index reports the same identities and text."""

    def test_native_units_match_the_shared_snapshot(self):
        snapshot = json.loads((FIXTURES / "native-units.json").read_text())["units"]
        for name, expected in snapshot.items():
            with self.subTest(name):
                self.assertEqual(deck_units((FIXTURES / name).read_bytes()), expected)


class SlideHighlightApiTest(unittest.TestCase):
    def setUp(self):
        from fastapi.testclient import TestClient

        from APP import db
        from APP.auth import get_current_user
        from APP.main import app

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        patcher = mock.patch.object(db, "DB_PATH", Path(tmp.name) / "test.db")
        patcher.start()
        self.addCleanup(patcher.stop)
        db.init_db()
        self.owner = db.create_user("slides@example.invalid")
        self.document = db.create_document(self.owner, "deck.pptx", pages=3)
        service = mock.Mock()
        service.get_document_pages.return_value = [{"page": p, "text": f"Slide {p}", "extraction_version": "1:abc"} for p in (1, 2, 3)]
        service_patch = mock.patch.object(app.state, "rag_service", service, create=True)
        service_patch.start()
        self.addCleanup(service_patch.stop)
        app.dependency_overrides[get_current_user] = lambda: {"id": self.owner}
        self.addCleanup(app.dependency_overrides.clear)
        self.client = TestClient(app)
        self.url = "/conversations/conv-p/annotations"
        self.payload = {"document_id": self.document, "kind": "manual-highlight", "page_number": 3, "quote": "171",
                        "extraction_version": "1:abc",
                        "anchor": {"format": "pptx", "slide_index": 2, "shape_id": 3, "row": 4, "col": 1, "quote": "171", "occurrence": 0}}

    def test_a_slide_highlight_round_trips_with_its_anchor(self):
        self.assertEqual(self.client.post(self.url, json=self.payload).status_code, 201)
        self.assertEqual(self.client.get(self.url).json()[0]["anchor"], self.payload["anchor"])

    def test_invalid_slide_anchors_are_refused(self):
        def anchor(**changes):
            return {**self.payload, "anchor": {**self.payload["anchor"], **changes}}
        for name, body, status in [
            ("wrong page for slide", {**self.payload, "page_number": 1}, 422),
            ("row without col", anchor(col=None), 422),
            ("empty quote", anchor(quote=""), 422),
            ("unknown key", anchor(range="A1"), 422),
            ("stale extraction", {**self.payload, "extraction_version": "1:old"}, 409),
        ]:
            with self.subTest(name):
                self.assertEqual(self.client.post(self.url, json=body).status_code, status)
        self.assertEqual(self.client.get(self.url).json(), [])


if __name__ == "__main__":
    unittest.main()
