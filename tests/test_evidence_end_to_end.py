"""One complete case: a PowerPoint with a picture on a slide.

Upload (the real ingest path) -> the picture is inventoried -> read by an OCR
provider through the spending ledger -> stored as its own evidence -> published
to both the keyword and the vector search -> retrieved -> cited, with the
slide's original text, citations and page version untouched throughout.

Only three things are fixtures: the OCR provider, the embedding model (a small
deterministic stand-in) and Qdrant (its in-memory mode)."""

import hashlib
import io
import re
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from qdrant_client import QdrantClient

from APP import db
from APP.extraction import store, worker
from APP.rag import evidence
from APP.rag import service as service_module
from APP.rag import vector_store
from APP.rag.service import RAGService

DIMENSIONS = 64


class Embeddings:
    """Bag of hashed words: texts sharing words are close. Deterministic."""

    def _vector(self, text):
        vector = [0.0] * DIMENSIONS
        for word in re.findall(r"[a-z0-9]+", text.lower()):
            vector[int(hashlib.sha256(word.encode()).hexdigest(), 16) % DIMENSIONS] += 1.0
        norm = sum(v * v for v in vector) ** 0.5 or 1.0
        return [v / norm for v in vector]

    def embed_documents(self, texts):
        return [self._vector(t) for t in texts]

    def embed_query(self, text):
        return self._vector(text)


class Ocr:
    version = "fake-ocr:test"

    def __init__(self):
        self.calls = 0

    def settings(self, mode):
        return {"version": self.version, "mode": mode}

    def ready(self):
        return True

    def estimate(self, mode):
        return 100

    def extract(self, data, mode):
        self.calls += 1
        return {"text": "Warranty claims fell 14 percent after the redesign", "provenance": "ocr", "content_type": "text",
                "extractor": self.version, "confidence": None, "pages": 1}


def deck_with_a_picture():
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[5])
    first.shapes.title.text = "Quarterly results overview"
    picture = io.BytesIO()
    Image.new("RGB", (300, 120), "white").save(picture, format="PNG")
    first.shapes.add_picture(io.BytesIO(picture.getvalue()), Inches(1), Inches(2), Inches(4))
    caption = first.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))    # text laid over the picture, as on most slides
    caption.text_frame.text = "Figure 1"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Hiring plan for next year"
    blank = deck.slides.add_slide(deck.slide_layouts[6])                                     # a slide with a picture and no text at all
    blank.shapes.add_picture(io.BytesIO(picture.getvalue()), Inches(1), Inches(1), Inches(3))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


class PowerPointWithAPictureTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.temp = Path(temp.name)
        for module in (db,):
            p = patch.object(module, "DB_PATH", self.temp / "test.db")
            p.start()
            self.addCleanup(p.stop)
        folder = lambda owner, document: self.temp / "docs" / owner / document
        for module in (service_module, vector_store):
            p = patch.object(module, "document_index_dir", folder)
            p.start()
            self.addCleanup(p.stop)
        p = patch.dict("os.environ", {"EVIDENCE_ENRICHMENT": "true", "NATIVE_CHUNKS_PPTX": "true", "QDRANT_COLLECTION": "e2e",
                                      "EXTRACTION_UPLOAD_CAP_MICROUSD": "10000", "EXTRACTION_DAILY_CAP_MICROUSD": "10000"})
        p.start()
        self.addCleanup(p.stop)
        db.init_db()
        self.owner = db.create_user("deck@example.invalid")
        self.service = RAGService.__new__(RAGService)
        self.service.settings = SimpleNamespace(qdrant_collection="e2e")
        self.service.qdrant = QdrantClient(":memory:")
        self.service.embedding_provider = Embeddings()
        self.service._embed_all = self.service.embedding_provider.embed_documents
        self.raw = deck_with_a_picture()
        stats = self.service.ingest(self.owner, "Deck.pptx", self.raw, 800, 100, 0.0)
        self.document = stats.document_id

    def chunks(self):
        _store, _bm25, chunks = self.service._load_document_artifacts(self.document, self.owner)
        return chunks

    def read_pictures(self, ocr=None):
        ocr = ocr or Ocr()
        while (job := store.claim()):
            worker.process(job, ocr)
        return ocr

    def test_the_picture_is_found_read_published_retrieved_and_cited_without_disturbing_the_slide_text(self):
        # 1. Upload: slide text indexed as before, the picture inventoried but not yet read.
        status = store.status(self.owner, self.document)
        self.assertEqual((status["total"], status["state"]), (2, "processing"))                  # two pictures, one on a slide with no text
        self.assertEqual(status["regions"][0]["location"]["format"], "pptx")
        self.assertEqual(status["regions"][0]["location"]["page"], 1)
        self.assertTrue(status["regions"][0]["location"]["overlaid"])                                # overlapped by the caption, still read
        self.assertEqual(status["regions"][0]["state"], "queued")
        shape_id = status["regions"][0]["location"]["shape_id"]
        self.assertEqual(store.revision(self.owner, self.document), 0)
        before = self.chunks()
        self.assertTrue(any("Quarterly results overview" in c.page_content for c in before))
        self.assertFalse(any("Warranty" in c.page_content for c in before))
        pages_file = self.service._page_texts_path(self.owner, self.document)
        pages_before = pages_file.read_bytes()
        native = next(c for c in before if "Quarterly results overview" in c.page_content)
        page = {"page": native.metadata["page"], "text": self.service._load_page_texts(self.owner, self.document)["pages"][str(native.metadata["page"])],
                "extraction_version": self.service._load_page_texts(self.owner, self.document)["extraction_version"]}
        cited_before = evidence.source_for(native, 1.0, page, None)
        key_before = self.service._summary_cache_key(self.owner, [self.document], "Summarize the deck")

        # 2. The picture is read once, through the ledger.
        ocr = self.read_pictures()
        self.assertEqual(ocr.calls, 1)
        with db._connect() as conn:
            charges = [tuple(r) for r in conn.execute("SELECT state, reserved, actual FROM extraction_charges")]
        self.assertEqual(charges, [("reconciled", 100, 100)])
        self.assertEqual(store.revision(self.owner, self.document), 0)                              # read, not yet searchable

        # 3. Published to both searches.
        self.service.publish_evidence(self.owner, self.document)
        self.assertEqual(store.revision(self.owner, self.document), 1)
        self.assertEqual(store.status(self.owner, self.document)["state"], "complete")
        vectorstore, bm25, after = self.service._load_document_artifacts(self.document, self.owner)
        added = [c for c in after if c.metadata.get("provenance") == "ocr"]
        self.assertEqual([c.page_content for c in added], ["Warranty claims fell 14 percent after the redesign"])   # the same picture twice: one chunk, one paid read
        self.assertEqual([place["page"] for place in added[0].metadata["also_at"]], [3])                           # and the other place it appears is kept
        records = [r for r in store.records(self.owner, self.document) if r["provenance"] == "ocr"]
        self.assertEqual(len(records), 2)                                                                          # both locations still have their own record
        self.assertEqual(bm25.corpus_size, len(after))                                               # the keyword index covers every chunk
        self.assertEqual(sum("warranty" in terms for terms in bm25.doc_freqs), 1)                    # and the picture text is in it, once
        hits = vectorstore.similarity_search("how much did warranty claims fall", k=3)
        self.assertEqual(hits[0].page_content, "Warranty claims fell 14 percent after the redesign")   # vector search finds it first
        self.assertTrue(any("Quarterly results overview" in c.page_content for c in after))        # the slide text is still there

        # 4. Cited as picture evidence, with the picture's own location.
        source = evidence.source_for(added[0], 1.0, page, None)
        self.assertEqual((source.provenance, source.page), ("ocr", 1))
        self.assertTrue(source.anchor["evidence_region"])
        self.assertEqual((source.anchor["slide_index"], source.anchor["shape_id"]), (0, shape_id))
        self.assertIsNone(source.page_char_start)                                                    # it never claims page-text offsets

        # 5. Nothing that existed moved.
        self.assertEqual(pages_file.read_bytes(), pages_before)
        cited_after = evidence.source_for(next(c for c in after if "Quarterly results overview" in c.page_content), 1.0, page, None)
        for field in ("citation_id", "page_char_start", "page_char_end", "extraction_version", "quote"):
            self.assertEqual(getattr(cited_after, field), getattr(cited_before, field), field)
        self.assertNotEqual(self.service._summary_cache_key(self.owner, [self.document], "Summarize the deck"), key_before)   # summaries refresh

    def test_every_citation_is_stamped_with_the_version_of_the_saved_pages_so_the_viewer_can_highlight_it(self):
        """A slide with no text still gets a page entry. The chunks must carry the version of
        the final page texts, or the viewer refuses every highlight ('precise passage location
        unavailable')."""
        payload = self.service._load_page_texts(self.owner, self.document)
        self.assertEqual(sorted(payload["pages"]), ["1", "2", "3"])                              # the textless slide has its page
        self.assertEqual(payload["pages"]["3"], "")
        for chunk in self.chunks():
            self.assertEqual(chunk.metadata["extraction_version"], payload["extraction_version"])
        native = next(c for c in self.chunks() if "Quarterly results overview" in c.page_content)
        source = evidence.source_for(native, 1.0, {"page": 1, "text": payload["pages"]["1"], "extraction_version": payload["extraction_version"]}, None)
        self.assertEqual(source.extraction_version, payload["extraction_version"])
        self.assertIsNotNone(source.page_char_start)                                              # located, not 'unavailable'
        self.assertIsNotNone(source.anchor)

    def test_reading_again_costs_nothing_and_a_second_publish_changes_nothing_but_the_revision(self):
        self.read_pictures()
        self.service.publish_evidence(self.owner, self.document)
        first = sorted(c.page_content for c in self.chunks())
        store.add_region(self.document, "v", {"format": "pptx", "page": 1, "slide_index": 0, "shape_id": 999}, asset_hash="same-image",
                         asset_path=str(self.temp / "x.png"), state="queued")
        with db._connect() as conn:
            hash_used = conn.execute("SELECT asset_hash FROM evidence_regions WHERE asset_hash IS NOT NULL ORDER BY created LIMIT 1").fetchone()[0]
            conn.execute("UPDATE evidence_regions SET asset_hash = ? WHERE asset_hash = 'same-image'", (hash_used,))
        (self.temp / "x.png").write_bytes(b"x")
        ocr = self.read_pictures()
        self.assertEqual(ocr.calls, 0)                                                               # the same image: served from the cache
        self.service.publish_evidence(self.owner, self.document)
        self.assertEqual(store.revision(self.owner, self.document), 1)                               # nothing new to search: no new revision
        self.assertEqual(sorted(c.page_content for c in self.chunks()), first)                       # and no duplicate chunk
        with db._connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM evidence_jobs WHERE state = 'indexed_pending'").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
