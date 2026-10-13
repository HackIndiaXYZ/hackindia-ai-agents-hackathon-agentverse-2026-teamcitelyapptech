"""Builds the synthetic XLSX and PPTX used by the office regression cases.

Every name, figure and sentence here is invented. The files are committed;
re-run this only to change them, then re-record the baseline:

    .venv/bin/python tests/fixtures/regression/make_fixtures.py
"""

from __future__ import annotations

import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent


def build_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.merge_cells("A1:D1")
    summary["A1"] = "Northwind Freight — FY2026 operating summary"
    summary["A1"].font = Font(bold=True, size=14)
    for col, text in zip("ABCD", ("Metric", "Actual", "Target", "Status")):
        summary[f"{col}3"] = text
        summary[f"{col}3"].font = Font(bold=True)
    rows = [
        ("Revenue", 48250000, 50000000, "Behind", "#,##0"),
        ("Gross margin", 0.412, 0.40, "Ahead", "0.0%"),
        ("Operating cash flow", 8400000, 7500000, "Ahead", "#,##0"),
        ("Net promoter score", 47, 55, "Behind", "0"),
        ("On-time delivery", 0.964, 0.95, "Ahead", "0.0%"),
    ]
    for i, (metric, actual, target, status, fmt) in enumerate(rows, start=4):
        summary[f"A{i}"] = metric
        summary[f"B{i}"] = actual
        summary[f"C{i}"] = target
        summary[f"D{i}"] = status
        summary[f"B{i}"].number_format = fmt
        summary[f"C{i}"].number_format = fmt
    summary["A10"] = "Report date"
    summary["B10"] = datetime.date(2026, 3, 31)
    summary["B10"].number_format = "yyyy-mm-dd"

    regions = wb.create_sheet("Regions")
    regions.append(["Region", "Q1", "Q2", "Q3", "Q4", "Regional lead"])
    region_rows = [
        ("Benelux", 2810000, 2955000, 3020000, 3180000, "Maarten de Vries"),
        ("Nordics", 3405000, 3560000, 3910000, 4025000, "Ingrid Solberg"),
        ("Iberia", 1980000, 2040000, 2115000, 2260000, "Lucía Ferrer"),
        ("DACH", 4120000, 4255000, 4390000, 4510000, "Jonas Albrecht"),
        ("UK & Ireland", 3770000, 3690000, 3820000, 3995000, "Siobhan Keane"),
        ("APAC South", 2230000, 2480000, 2610000, 2875000, "Arjun Mehta"),
        ("APAC North", 2655000, 2700000, 2790000, 2840000, "Mei Tanaka"),
        ("Middle East", 1540000, 1620000, 1755000, 1890000, "Omar Haddad"),
    ]
    for row in region_rows:
        regions.append(list(row))
    for r in range(2, 2 + len(region_rows)):
        for c in "BCDE":
            regions[f"{c}{r}"].number_format = "#,##0"

    headcount = wb.create_sheet("Headcount")
    headcount.append(["Department", "Location", "FTE", "Open roles"])
    for row in [
        ("Engineering", "Pune", 84, 12),
        ("Engineering", "Rotterdam", 41, 3),
        ("Operations", "Rotterdam", 132, 9),
        ("Operations", "Gdańsk", 97, 14),
        ("Finance", "Rotterdam", 23, 1),
        ("Customer success", "Lisbon", 38, 6),
    ]:
        headcount.append(list(row))

    risks = wb.create_sheet("Risks")
    risks.append(["ID", "Risk", "Owner", "Mitigation"])
    for row in [
        ("R1", "Supplier concentration in trailer leasing",
         "Petra Novak", "Qualify a second leasing partner by Q3 and cap any single lessor at 45% of the fleet."),
        ("R2", "Data residency for EU customer records",
         "Daniel Okafor", "Move the remaining customer records to the Frankfurt region before the October audit."),
        ("R3", "Fuel price volatility",
         "Lucía Ferrer", "Hedge 60% of second-half diesel volume through fixed-price contracts."),
        ("R4", "Driver attrition in the Nordics",
         "Ingrid Solberg", "Introduce a retention bonus after 18 months of service."),
    ]:
        risks.append(list(row))

    monthly = wb.create_sheet("Monthly")
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    monthly.append(["Product", *months, "Total"])
    for name, base, step in [("Starter", 410, 7), ("Professional", 1220, 15), ("Enterprise", 318, 4)]:
        values = [base + step * i for i in range(12)]
        monthly.append([name, *values, sum(values)])

    wb.save(path)


def build_pptx(path: Path) -> None:
    deck = Presentation()
    blank_title = deck.slide_layouts[5]

    s = deck.slides.add_slide(deck.slide_layouts[0])
    s.shapes.title.text = "Atlas Logistics — Board update"
    s.placeholders[1].text = "Second quarter, FY2026"

    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = "Quarter highlights"
    body = s.placeholders[1].text_frame
    body.text = "On-time delivery reached 96.4%, up from 93.8% a year ago."
    for line in ("The Rotterdam cross-dock opened in May and now handles 1,150 pallets a day.",
                 "Customer churn fell to 3.1% of annual recurring revenue.",
                 "Safety remains our first priority."):
        body.add_paragraph().text = line

    s = deck.slides.add_slide(blank_title)
    s.shapes.title.text = "Fleet by region"
    table = s.shapes.add_table(5, 3, Inches(1), Inches(1.8), Inches(8), Inches(2.5)).table
    for r, row in enumerate([("Region", "Trucks", "Utilisation"),
                             ("Benelux", "140", "87%"),
                             ("Nordics", "96", "81%"),
                             ("Iberia", "58", "74%"),
                             ("DACH", "171", "90%")]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text

    s = deck.slides.add_slide(blank_title)
    s.shapes.title.text = "Principal risks"
    box = s.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(1.2))
    box.name = "Risk statement"
    box.text_frame.word_wrap = True
    box.text_frame.text = ("Fuel price volatility remains the top risk; hedging covers 60% of "
                           "second-half diesel volume.")
    box2 = s.shapes.add_textbox(Inches(1), Inches(3.4), Inches(8), Inches(1.2))
    box2.name = "Second risk"
    box2.text_frame.word_wrap = True
    box2.text_frame.text = "Driver attrition in the Nordics is 19%, against a group average of 11%."

    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = "Hiring plan"
    body = s.placeholders[1].text_frame
    body.text = "Add 45 drivers in the third quarter, mostly in Benelux and DACH."
    body.add_paragraph().text = "Hire a head of data protection before the October audit."

    s = deck.slides.add_slide(blank_title)
    s.shapes.title.text = "Capital expenditure"
    left = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4), Inches(1.5))
    left.name = "Capex budget"
    left.text_frame.text = "Capex budget for FY2026: €12.5 million."
    right = s.shapes.add_textbox(Inches(5.2), Inches(1.8), Inches(4), Inches(1.5))
    right.name = "Capex spent"
    right.text_frame.text = "Spent to date: €5.2 million, mainly electric yard tractors."

    s = deck.slides.add_slide(blank_title)
    s.shapes.title.text = "Safety"
    box = s.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(1.5))
    box.name = "Safety statement"
    box.text_frame.word_wrap = True
    box.text_frame.text = "Safety remains our first priority."
    box.text_frame.add_paragraph().text = "Lost-time injuries fell to 0.8 per million hours worked."

    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = "Next steps"
    body = s.placeholders[1].text_frame
    body.text = "Approve the second leasing partner at the September board."
    body.add_paragraph().text = "Report hedging results at the year-end review."
    for p in body.paragraphs:
        for run in p.runs:
            run.font.size = Pt(24)

    deck.save(path)


if __name__ == "__main__":
    build_xlsx(HERE / "northwind_operations.xlsx")
    build_pptx(HERE / "atlas_board_update.pptx")
    print("written:", *(p.name for p in sorted(HERE.glob("*.xlsx")) + sorted(HERE.glob("*.pptx"))))
